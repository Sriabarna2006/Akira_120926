import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.pdfgen import canvas

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def create_docx(filename):
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("AKIRA — Comprehensive Project Architecture & Evolution")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(10, 37, 64)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Full Process, Architecture, Engineering Specs & Working Workflows (Phases 1 to 4)")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(70, 80, 95)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run("System: AKIRA (Autonomous Knowledge & Intelligence Real-time Aggregator)\nStatus: Phases 1, 2, 3, 4 Completed & Verified | Date: September 2026")
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(100, 110, 120)
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("―" * 55)

    def add_h1(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = RGBColor(14, 116, 144) # Cyan/Dark Teal
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)
        return h

    def add_h3(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(71, 85, 105)
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)
        return h

    def add_body(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        return p

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        r_bold = p.add_run(bold_prefix + ": ")
        r_bold.font.bold = True
        r_bold.font.size = Pt(10)
        r_text = p.add_run(text)
        r_text.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        return p

    # Section 1: Executive Overview
    add_h1("1. EXECUTIVE PROJECT OVERVIEW & PHILOSOPHY")
    add_body("AKIRA (Autonomous Knowledge & Intelligence Real-time Aggregator) is a next-generation real-world intelligence and news aggregation platform. Unlike conventional news scrapers or click-driven media apps, AKIRA is engineered around factual corroboration, strict source transparency, anti-hallucination pipelines, and canonical event clustering.")
    add_bullet("Core Principle 1 — Event vs Article Separation", "An Article is a specific report from a publisher. A Canonical Event is the underlying real-world occurrence. Multiple publisher reports are clustered into one single Canonical Event, preventing duplicate clutter while providing multi-source corroboration.")
    add_bullet("Core Principle 2 — First-Class Regional Focus", "Tamil Nadu is treated as a top-tier primary region alongside National (India) and Global (World), with dedicated district-level, infrastructure, government (DIPR), and technological taxonomy.")
    add_bullet("Core Principle 3 — No Fake Real-Time Data", "AKIRA enforces transparent timestamps (published_at vs discovered_at) and never displays deceptive 'BREAKING' or 'LIVE' flags unless backed by verifiable data.")
    add_bullet("Core Principle 4 — Strict Source Transparency", "Every article preserves its original publisher identity, credibility tier, and unaltered destination URL.")

    # Section 2: Phase 1
    add_h1("2. PHASE 1: DATABASE & AUTHENTICATION FOUNDATIONS")
    add_h2("A. Normal Working Phase (User & Product Experience)")
    add_body("Phase 1 established the foundation for user accounts, security, and private data isolation. Users can register with email/password, log in securely, obtain authenticated JWT sessions, manage profile settings, and bookmark saved stories. A user's bookmarks are completely private and cannot be viewed or accessed by other users.")
    
    add_h2("B. Technical Engineering Phase (Architecture & Code)")
    add_bullet("PostgreSQL / Neon / Supabase Schema", "Designed relational foundation with UUID keys, foreign key constraints, timestamps, and indexes across core tables (profiles, saved_events, saved_articles).")
    add_bullet("Row-Level Security (RLS)", "Configured Postgres RLS policies ensuring users can only read and mutate their own saved events and profile records.")
    add_bullet("JWT Authentication Middleware", "Implemented Bearer token verification using Supabase Auth JWTs, extracting user claims and roles.")
    add_bullet("Fallback Architecture", "Engineered in-memory repository fallbacks allowing full local offline development without crashing when remote credentials are missing.")
    add_bullet("Migration Management", "Created versioned SQL migrations (001_initial_schema.sql, 002_rls_policies.sql, 003_seed_data.sql).")

    # Section 3: Phase 2
    add_h1("3. PHASE 2: FRONTEND MVP & UI FOUNDATION")
    add_h2("A. Normal Working Phase (User & Product Experience)")
    add_body("Phase 2 delivered a world-class, responsive, dark-mode user interface designed to wow users. It features sleek glassmorphism panels, cyan/indigo accent glows, smooth micro-interactions, responsive sidebars, intuitive regional filters (Tamil Nadu, India, World), category badges, interactive live news streams, search drawers, bookmark management, and rich event detail views.")

    add_h2("B. Technical Engineering Phase (Architecture & Code)")
    add_bullet("Technology Stack", "React 18, TypeScript, Vite, React Router v6, Lucide React icons, TanStack React Query.")
    add_bullet("Custom CSS Design System", "Strictly Tailwind-free, handcrafted vanilla CSS design system (index.css) utilizing CSS Custom Properties, backdrop filters, glass tokens, smooth cubic-bezier transitions, and accessible font hierarchies.")
    add_bullet("Modular Component Hierarchy", "Engineered decoupled component architecture across layout/ (Navbar, Sidebar, Footer, MobileNav), events/ (EventCard, SourceBadge, EventTimeline, UrgencyTag), auth/ (AuthModal, UserDropdown), and common/ (LoadingSpinner, ErrorState, EmptyState).")
    add_bullet("Client Route Architecture", "Implemented routing for Home/Explore (/), Live Trending (/live), Regional Feeds (/region/:id), Category Domains (/category/:id), Search (/search), Bookmarks (/saved), and Event Detail (/event/:id).")

    # Section 4: Phase 3
    add_h1("4. PHASE 3: CORE DATA LAYER & BACKEND FOUNDATION")
    add_h2("A. Normal Working Phase (User & Product Experience)")
    add_body("Phase 3 connected the frontend to a high-performance Express REST API and live PostgreSQL database. Users now browse real persisted categories (15 knowledge domains), explore vetted news publishers, paginate through historical articles, filter canonical events by urgency and geography, and perform keyword searches.")

    add_h2("B. Technical Engineering Phase (Architecture & Code)")
    add_bullet("Layered Backend Architecture", "Routes -> Middleware (Auth, RateLimit, Validator) -> Controllers -> Services -> Repositories -> PostgreSQL Database.")
    add_bullet("Zod Schema Validation", "Constructed comprehensive query and parameter schemas (query.validator.ts), rejecting malformed inputs with standardized 400 Bad Request responses.")
    add_bullet("Sliding Window Rate Limiting", "Implemented sliding window memory rate limiters preventing abuse and brute-force flooding.")
    add_bullet("Safe Content Sanitizer", "Engineered strict HTML sanitizer (sanitize.ts) that strips dangerous tags (<script>, <iframe>, event attributes) and enforces safe HTTP/HTTPS URL protocols.")
    add_bullet("Protected Ingestion Sync", "Created POST /api/live/sync guarded with internal secret header (x-akira-internal-key) and admin verification.")

    # Section 5: Phase 4
    add_h1("5. PHASE 4: REAL NEWS INGESTION & EVENT PIPELINE")
    add_h2("A. Normal Working Phase (User & Product Experience)")
    add_body("Phase 4 activated AKIRA's real-world information gathering engine. AKIRA automatically pulls real news from 14 verified international, national, and Tamil Nadu publishers. Incoming news reports are automatically checked for duplicates, classified into the right region and topic, and intelligently merged into Canonical Events when multiple publishers report the same story. Users see a clean, multi-corroborated news feed with transparent publisher badges, publication times, and direct source links.")

    add_h2("B. Technical Engineering Phase (Architecture & Code)")
    add_bullet("Network Fetching & SSRF Safety", "Engineered feedFetcher.ts using rss-parser with Chrome User-Agent, SSRF URL whitelist check, strict 6,000ms Promise.race timeout, and exponential retry backoff.")
    add_bullet("URL Normalization Engine", "Engineered urlNormalizer.ts to strip tracking parameters (utm_*, ref, fbclid, gclid, hash fragments) while preserving destination URLs.")
    add_bullet("Article Validation & Deduplication", "articleValidator.ts validates title length (5-500 chars), future timestamp bounds, and source identity. article.repository.ts checks existsByUrl for instant idempotent skipping.")
    add_bullet("Deterministic Layered Classifier", "classifier.ts assigns regions with Tamil Nadu priority (districts, Chennai metro, CM, DIPR) over India/World, and categorizes into 15 standardized knowledge domains (Region != Category).")
    add_bullet("Canonical Event Clustering Engine", "eventMatcher.ts clusters reports using 36h temporal window, regional isolation (preventing false merging of NASA vs ISRO missions), and title token Jaccard similarity (0.60 title + 0.40 body, threshold 0.28).")
    add_bullet("Idempotency & Multi-Source Attachment", "attachArticleToEvent updates event_sources (ON CONFLICT DO NOTHING), links article to event_id, and refreshes last_updated_at. Consecutive runs result in 100% duplicate skipping (0 duplicates created).")
    add_bullet("Background Scheduling & Health", "scheduler.ts runs ingestion every 15 minutes with re-entrancy locks (isSyncing), tracking last_successful_fetch, failure_count, and healthy/degraded status per source.")

    # Section 6: Comprehensive Architecture Diagram
    add_h1("6. END-TO-END DATA FLOW DIAGRAM")
    add_body("The complete AKIRA pipeline operates as follows:")
    
    flow_text = (
        "14 Approved RSS/Atom Feeds (The Hindu, Express, BBC, Reuters, TechCrunch, ScienceDaily)\n"
        "  ↓\n"
        "Feed Fetcher (User-Agent, SSRF Guard, 6s Timeout, Exponential Backoff)\n"
        "  ↓\n"
        "URL Normalization (Strip UTM, ref, fbclid) & HTML Sanitization (Strip scripts)\n"
        "  ↓\n"
        "Validation (Title length, safe protocol, timestamp bounds)\n"
        "  ↓\n"
        "Deduplication (URL existence check in DB)\n"
        "  ↓\n"
        "Deterministic Classification (Tamil Nadu priority, India, World; 15 Categories)\n"
        "  ↓\n"
        "Canonical Event Clustering (36h temporal window, Regional isolation, Title Jaccard)\n"
        "  ↓\n"
        "PostgreSQL Persistence (articles, canonical_events, event_sources)\n"
        "  ↓\n"
        "REST API Controllers & Services (with RateLimit, Zod Validator, JWT Auth)\n"
        "  ↓\n"
        "React 18 Frontend (Live Trending, Explore, Category, Event Detail, Bookmarks)"
    )
    add_body(flow_text)

    # Section 7: Complete File Inventory
    add_h1("7. COMPLETE FILE & MODULE INVENTORY ADDED TO DATE")
    
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Module / Layer"
    hdr_cells[1].text = "Key File Paths"
    hdr_cells[2].text = "Core Functionality"
    
    for c in hdr_cells:
        set_cell_background(c, "0E7490")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9.5)

    modules_data = [
        ("Database & Migrations", "server/src/db/migrations/001_initial_schema.sql\n002_rls_policies.sql\n003_seed_data.sql\nneon_schema.sql", "Postgres tables, foreign keys, RLS security policies, reference data for 14 feeds and 15 categories."),
        ("Ingestion Pipeline", "server/src/services/ingestion/\n- feedFetcher.ts\n- urlNormalizer.ts\n- articleValidator.ts\n- classifier.ts\n- eventMatcher.ts\n- scheduler.ts", "RSS/Atom fetching, timeout handling, tracking stripper, HTML sanitizer, classification, clustering, cron scheduler."),
        ("Repositories", "server/src/repositories/\n- region.repository.ts\n- category.repository.ts\n- source.repository.ts\n- article.repository.ts\n- event.repository.ts", "PostgreSQL database access layer with in-memory fallbacks, conflict resolution, and relational joins."),
        ("Services & Controllers", "server/src/services/newsIngestion.service.ts\nserver/src/controllers/\n- live.controller.ts\n- news.controller.ts\n- region.controller.ts\n- category.controller.ts", "Business logic orchestration, API endpoints for /api/live, /api/events, /api/sources, and protected /api/live/sync."),
        ("Security & Middleware", "server/src/middleware/\n- auth.middleware.ts\n- rateLimit.middleware.ts\n- validate.middleware.ts\n- errorHandler.ts\nserver/src/utils/sanitize.ts", "Supabase JWT verification, sliding window rate limiter, Zod request validator, centralized error logger, XSS sanitizer."),
        ("Frontend Pages", "client/src/pages/\n- HomePage.tsx\n- LiveTrendingPage.tsx\n- ExplorePage.tsx\n- CategoryPage.tsx\n- EventDetailPage.tsx\n- AllNewsPage.tsx\n- SavedPage.tsx", "Responsive React pages with glassmorphism UI, real-time data hooks, multi-source corroboration badges, and search."),
        ("Frontend Components", "client/src/components/\n- layout/ (Navbar, Sidebar, MobileNav)\n- events/ (EventCard, SourceBadge)\n- auth/ (AuthModal)\n- common/ (LoadingSpinner, ErrorState)", "Reusable UI components, accessible navigation, modal dialogs, and responsive layout drawers."),
        ("Automated Tests", "server/src/tests/\n- phase1.test.ts\n- phase3.test.ts\n- phase4.test.ts", "Comprehensive automated verification suites testing auth, RLS, REST APIs, feed parsing, deduplication, and idempotency.")
    ]

    for mod, files, desc in modules_data:
        row_cells = table.add_row().cells
        row_cells[0].text = mod
        row_cells[1].text = files
        row_cells[2].text = desc
        for i, c in enumerate(row_cells):
            set_cell_background(c, "F8FAFC" if len(table.rows) % 2 == 0 else "FFFFFF")
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)

    # Section 8: Verification & Status
    add_h1("8. VERIFICATION RESULTS & SYSTEM STATUS")
    add_body("All four project phases have been fully verified with automated test suites and production builds:")
    add_bullet("Phase 1 Auth & DB Suite", "PASS (7/7 tests passed: Migrations, JWT Auth, RLS Private Isolation, Sync Guard)")
    add_bullet("Phase 3 Core Data API Suite", "PASS (14/14 tests passed: Repositories, REST Endpoints, Rate Limiting, Zod Validation, Sanitization)")
    add_bullet("Phase 4 News Ingestion Suite", "PASS (9/9 tests passed: URL Normalization, HTML Sanitizer, Validation, Classifier, Clustering, RSS Fetch, 100% Idempotency, Health Tracking, Protected Sync)")
    add_bullet("Client Production Build", "PASS (Vite build completed: 1,783 modules transformed, 0 errors)")
    add_bullet("Server TypeScript Build", "PASS (tsc compilation completed: 0 errors)")

    doc.save(filename)
    print(f"Successfully generated DOCX: {filename}")

def create_pdf(filename):
    doc = SimpleDocTemplate(
        filename,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0F172A'),
        alignment=1, # Center
        spaceAfter=4
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#0E7490'),
        alignment=1,
        spaceAfter=6
    )
    
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#64748B'),
        alignment=1,
        spaceAfter=12
    )

    h1_style = ParagraphStyle(
        'CustomH1',
        parent=styles['Heading1'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#0E7490'),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'CustomH2',
        parent=styles['Heading2'],
        fontSize=10.5,
        leading=13,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=7,
        spaceAfter=3,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['Normal'],
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=4
    )

    bullet_style = ParagraphStyle(
        'CustomBullet',
        parent=styles['Normal'],
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#334155'),
        leftIndent=12,
        spaceAfter=2
    )

    story = []

    story.append(Paragraph("AKIRA — Complete Project Architecture & Evolution", title_style))
    story.append(Paragraph("Full Technical Specifications, Architecture & Operational Workflows (Phases 1 to 4)", subtitle_style))
    story.append(Paragraph("System: AKIRA Intelligence Platform | Date: September 2026 | Status: Verified & Approved", meta_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=8))

    # Section 1
    story.append(Paragraph("1. EXECUTIVE PROJECT OVERVIEW & PHILOSOPHY", h1_style))
    story.append(Paragraph("AKIRA (Autonomous Knowledge & Intelligence Real-time Aggregator) is a real-world intelligence aggregation platform built on fact corroboration, source transparency, anti-hallucination pipelines, and canonical event clustering.", body_style))
    story.append(Paragraph("• <b>Canonical Event vs Article Model:</b> Articles are publisher reports. Canonical Events are real-world occurrences. Multiple reports merge into 1 canonical event with corroborating source badges.", bullet_style))
    story.append(Paragraph("• <b>First-Class Regional Priority:</b> Tamil Nadu is treated as a top-tier primary region with dedicated infrastructure, public announcement, and district taxonomies alongside India and World.", bullet_style))
    story.append(Paragraph("• <b>Anti-Hallucination & No Fake Live Data:</b> Transparent timestamps (published_at vs discovered_at) ensure no fabricated 'BREAKING' or 'LIVE' tags.", bullet_style))

    # Section 2
    story.append(Paragraph("2. PHASE 1: DATABASE & AUTHENTICATION FOUNDATIONS", h1_style))
    story.append(Paragraph("<b>Normal Working Phase:</b> Users can register, log in with JWT sessions, manage profiles, and save bookmarked events. User saved data is strictly private and isolated.", body_style))
    story.append(Paragraph("<b>Technical Phase:</b> PostgreSQL/Supabase schema with UUID keys, Row-Level Security (RLS) policies, JWT middleware, in-memory repository fallbacks, and automated test migrations.", body_style))

    # Section 3
    story.append(Paragraph("3. PHASE 2: FRONTEND MVP & UI FOUNDATION", h1_style))
    story.append(Paragraph("<b>Normal Working Phase:</b> High-impact responsive dark UI featuring glassmorphism cards, cyan/indigo accent lighting, region selector tabs, category pills, live stream explorer, search drawers, and bookmark views.", body_style))
    story.append(Paragraph("<b>Technical Phase:</b> React 18, TypeScript, Vite, React Router v6, Lucide icons, TanStack React Query, and a handcrafted 100% Tailwind-free Vanilla CSS design system (index.css).", body_style))

    # Section 4
    story.append(Paragraph("4. PHASE 3: CORE DATA LAYER & BACKEND FOUNDATION", h1_style))
    story.append(Paragraph("<b>Normal Working Phase:</b> Connects UI to live Express REST API and PostgreSQL database, supporting pagination, region/category filters, search, and canonical event aggregation.", body_style))
    story.append(Paragraph("<b>Technical Phase:</b> Layered Express architecture (Routes -> Middleware -> Controllers -> Services -> Repositories -> Database), Zod request validation, sliding window rate limiter, safe HTML sanitizer, and internal key authentication.", body_style))

    # Section 5
    story.append(Paragraph("5. PHASE 4: REAL NEWS INGESTION & EVENT PIPELINE", h1_style))
    story.append(Paragraph("<b>Normal Working Phase:</b> Ingests live news from 14 vetted publishers (The Hindu, Indian Express, BBC, Reuters, TechCrunch, ScienceDaily, etc.). Automatically sanitizes, validates, deduplicates, classifies regions (Tamil Nadu priority), and clusters corroborating reports into single Canonical Events with transparent source badges.", body_style))
    story.append(Paragraph("<b>Technical Phase:</b> rss-parser with browser User-Agent, SSRF whitelist guard, 6,000ms timeout with exponential retry, tracking parameter stripper (urlNormalizer.ts), deterministic classifier (classifier.ts), 36h temporal & title token Jaccard event matcher (eventMatcher.ts), 100% idempotent PostgreSQL persistence, and background cron scheduler (scheduler.ts).", body_style))

    story.append(Spacer(1, 4))
    story.append(Paragraph("6. COMPREHENSIVE FILE & MODULE INVENTORY", h1_style))

    table_data = [
        [Paragraph("<b>Layer / Module</b>", body_style), Paragraph("<b>Key File Paths</b>", body_style), Paragraph("<b>Core Responsibility</b>", body_style)],
        [Paragraph("Database & Seed", body_style), Paragraph("server/src/db/migrations/<br/>001_initial_schema.sql<br/>002_rls_policies.sql<br/>003_seed_data.sql<br/>neon_schema.sql", body_style), Paragraph("Core tables, RLS policies, foreign keys, and seed data for 14 sources & 15 categories.", body_style)],
        [Paragraph("Ingestion Pipeline", body_style), Paragraph("server/src/services/ingestion/<br/>feedFetcher.ts, urlNormalizer.ts,<br/>articleValidator.ts, classifier.ts,<br/>eventMatcher.ts, scheduler.ts", body_style), Paragraph("Network fetching, timeout/retry, URL cleaning, HTML sanitizer, classification, clustering, background cron.", body_style)],
        [Paragraph("Repositories", body_style), Paragraph("server/src/repositories/<br/>source, article, event, region, category .repository.ts", body_style), Paragraph("PostgreSQL queries, in-memory fallbacks, URL existence deduplication, and event-source attachment.", body_style)],
        [Paragraph("API Controllers & Middleware", body_style), Paragraph("server/src/controllers/ & server/src/middleware/<br/>auth, rateLimit, validate, errorHandler", body_style), Paragraph("REST API endpoints, JWT auth, Zod query validation, sliding rate limiting, error logging.", body_style)],
        [Paragraph("Frontend UI", body_style), Paragraph("client/src/pages/ & client/src/components/<br/>LiveTrending, Explore, EventDetail, AllNews, index.css", body_style), Paragraph("React 18 pages, glassmorphism design system, multi-source corroboration badges, and search.", body_style)],
        [Paragraph("Automated Tests", body_style), Paragraph("server/src/tests/<br/>phase1.test.ts, phase3.test.ts, phase4.test.ts", body_style), Paragraph("Automated verification for Auth, RLS, REST APIs, Ingestion, Deduplication, and 100% Idempotency.", body_style)]
    ]

    t = Table(table_data, colWidths=[1.3*inch, 2.3*inch, 3.4*inch])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0E7490')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F8FAFC'), colors.white]),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
    ]))
    story.append(t)

    story.append(Spacer(1, 6))
    story.append(Paragraph("7. VERIFICATION RESULTS & SYSTEM STATUS", h1_style))
    story.append(Paragraph("• <b>Phase 1 Test Suite:</b> 7/7 PASSED (Database migrations, JWT auth, private RLS isolation, sync guard)", bullet_style))
    story.append(Paragraph("• <b>Phase 3 Test Suite:</b> 14/14 PASSED (Repositories, REST APIs, Zod validation, XSS sanitization, rate limits)", bullet_style))
    story.append(Paragraph("• <b>Phase 4 Ingestion Suite:</b> 9/9 PASSED (URL normalization, validation, classifier, event clustering, 100% idempotency, health tracking)", bullet_style))
    story.append(Paragraph("• <b>Build Verification:</b> Client Vite production build PASSED (1,783 modules, 0 errors); Server TypeScript build PASSED (0 errors).", bullet_style))

    doc.build(story)
    print(f"Successfully generated PDF: {filename}")

if __name__ == '__main__':
    workspace_dir = r"a:\My web app(ML ai)"
    docx_path = os.path.join(workspace_dir, "AKIRA_Complete_Project_Documentation.docx")
    pdf_path = os.path.join(workspace_dir, "AKIRA_Complete_Project_Documentation.pdf")
    
    create_docx(docx_path)
    create_pdf(pdf_path)
